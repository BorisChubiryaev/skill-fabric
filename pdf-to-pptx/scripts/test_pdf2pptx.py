"""Юнит-тесты инструмента pdf2pptx.

Запуск:  python3 -m unittest scripts/test_pdf2pptx.py -v
Требуются pymupdf и python-pptx.
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
TOOL = os.path.join(HERE, "pdf2pptx.py")
sys.path.insert(0, HERE)

import pymupdf as fitz  # noqa: E402
from pptx import Presentation  # noqa: E402


def _run(*args: str) -> tuple[int, dict]:
    proc = subprocess.run(
        [sys.executable, TOOL, *args], capture_output=True, text=True
    )
    try:
        payload = json.loads(proc.stdout)
    except json.JSONDecodeError:
        payload = {"_stdout": proc.stdout, "_stderr": proc.stderr}
    return proc.returncode, payload


def _make_deck(path: str, pages: int = 2, with_text: bool = True) -> None:
    doc = fitz.open()
    for i in range(pages):
        p = doc.new_page(width=720, height=540)
        p.draw_rect(fitz.Rect(0, 0, 720, 100), fill=(0.1, 0.3, 0.6))
        p.draw_circle(fitz.Point(600, 300), 50, fill=(0.9, 0.4, 0.1))
        if with_text:
            p.insert_text((40, 70), f"Заголовок {i}", fontsize=30, color=(1, 1, 1))
            p.insert_text((40, 160), "Обычный текст строкой", fontsize=16)
    doc.save(path)
    doc.close()


class ToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp()
        self.pdf = os.path.join(self.tmp, "deck.pdf")
        _make_deck(self.pdf, pages=2)

    def test_analyze_reports_pages_and_text(self) -> None:
        rc, out = _run("analyze", self.pdf)
        self.assertEqual(rc, 0)
        self.assertTrue(out["ok"])
        self.assertEqual(out["page_count"], 2)
        self.assertGreater(out["total_chars"], 0)
        self.assertEqual(out["scanned_pages"], 0)
        self.assertEqual(out["recommended_mode"], "hybrid")

    def test_analyze_detects_scan(self) -> None:
        scan = os.path.join(self.tmp, "scan.pdf")
        _make_deck(scan, pages=1, with_text=False)
        rc, out = _run("analyze", scan)
        self.assertEqual(rc, 0)
        self.assertEqual(out["total_chars"], 0)
        self.assertEqual(out["scanned_pages"], 1)
        self.assertEqual(out["recommended_mode"], "image")

    def test_convert_hybrid_produces_editable_text_and_backgrounds(self) -> None:
        out_pptx = os.path.join(self.tmp, "out.pptx")
        report = os.path.join(self.tmp, "out.md")
        rc, out = _run("convert", self.pdf, "--out", out_pptx,
                       "--report", report, "--mode", "hybrid")
        self.assertEqual(rc, 0)
        self.assertTrue(out["ok"])
        self.assertEqual(out["stats"]["slides"], 2)
        self.assertEqual(out["stats"]["backgrounds"], 2)
        self.assertGreater(out["stats"]["text_chars"], 0)
        self.assertTrue(os.path.getsize(out_pptx) > 0)
        self.assertTrue(os.path.getsize(report) > 0)
        # Текст действительно редактируемый: есть текстовые фигуры с символами.
        prs = Presentation(out_pptx)
        chars = sum(
            len(sh.text_frame.text)
            for sl in prs.slides for sh in sl.shapes if sh.has_text_frame
        )
        self.assertGreater(chars, 0)

    def test_convert_text_mode_drops_vector_but_keeps_text(self) -> None:
        out_pptx = os.path.join(self.tmp, "t.pptx")
        rc, out = _run("convert", self.pdf, "--out", out_pptx, "--mode", "text")
        self.assertEqual(rc, 0)
        self.assertEqual(out["stats"]["backgrounds"], 0)
        self.assertGreater(out["stats"]["text_chars"], 0)

    def test_convert_vector_mode_keeps_text_and_editable_shapes(self) -> None:
        out_pptx = os.path.join(self.tmp, "vec.pptx")
        rc, out = _run("convert", self.pdf, "--out", out_pptx, "--mode", "vector")
        self.assertEqual(rc, 0)
        # Никакой растровой подложки: графика — редактируемые фигуры.
        self.assertEqual(out["stats"]["backgrounds"], 0)
        self.assertGreater(out["stats"]["vector_shapes"], 0)
        self.assertGreater(out["stats"]["text_chars"], 0)
        # В PPTX не должно быть картинок, но должны быть фигуры-фриформы.
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(out_pptx)
        pictures = freeforms = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures += 1
                if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM:
                    freeforms += 1
        self.assertEqual(pictures, 0)
        self.assertGreater(freeforms, 0)

    def test_convert_image_mode_has_no_text(self) -> None:
        out_pptx = os.path.join(self.tmp, "i.pptx")
        rc, out = _run("convert", self.pdf, "--out", out_pptx, "--mode", "image")
        self.assertEqual(rc, 0)
        self.assertEqual(out["stats"]["text_chars"], 0)
        self.assertEqual(out["stats"]["backgrounds"], 2)

    def test_refuses_overwriting_input(self) -> None:
        rc, out = _run("convert", self.pdf, "--out", self.pdf, "--mode", "hybrid")
        self.assertEqual(rc, 1)
        self.assertFalse(out["ok"])

    def test_refuses_non_pptx_out(self) -> None:
        rc, out = _run("convert", self.pdf, "--out",
                       os.path.join(self.tmp, "x.ppt"), "--mode", "hybrid")
        self.assertEqual(rc, 1)
        self.assertFalse(out["ok"])

    def test_verify_counts_slides_and_shapes(self) -> None:
        out_pptx = os.path.join(self.tmp, "v.pptx")
        _run("convert", self.pdf, "--out", out_pptx, "--mode", "hybrid")
        rc, out = _run("verify", out_pptx)
        self.assertEqual(rc, 0)
        self.assertTrue(out["ok"])
        self.assertEqual(out["slides"], 2)
        self.assertEqual(out["slide_width_pt"], 720.0)
        self.assertEqual(out["slide_height_pt"], 540.0)
        self.assertGreater(out["total_shapes"], 0)

    def test_missing_input(self) -> None:
        rc, out = _run("analyze", os.path.join(self.tmp, "nope.pdf"))
        self.assertEqual(rc, 1)
        self.assertFalse(out["ok"])

    def test_color_and_font_helpers(self) -> None:
        import pdf2pptx as m
        rgb = m._int_color_to_rgb(0x0033AA)
        self.assertEqual((rgb[0], rgb[1], rgb[2]), (0x00, 0x33, 0xAA))
        self.assertIsNone(m._int_color_to_rgb(None))
        self.assertEqual(m._clean_font_name("ABCDEF+Arial-BoldMT"), "Arial")

    def test_float_color_and_bezier_helpers(self) -> None:
        import pdf2pptx as m
        rgb = m._float_color_to_rgb((1.0, 0.0, 0.5))
        self.assertEqual((rgb[0], rgb[1], rgb[2]), (0xFF, 0x00, 0x80))
        # Одноканальный (серый) цвет разворачивается в RGB.
        grey = m._float_color_to_rgb((0.5,))
        self.assertEqual((grey[0], grey[1], grey[2]), (0x80, 0x80, 0x80))
        self.assertIsNone(m._float_color_to_rgb(None))
        p = fitz.Point
        flat = m._flatten_cubic(p(0, 0), p(0, 10), p(10, 10), p(10, 0), steps=4)
        self.assertEqual(len(flat), 4)
        self.assertEqual(flat[-1], (10.0, 0.0))


if __name__ == "__main__":
    unittest.main()
